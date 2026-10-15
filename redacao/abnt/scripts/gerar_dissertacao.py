#!/usr/bin/env python3
"""Gera a dissertacao em formato ABNT (DOCX) a partir dos capitulos em LaTeX (somente leitura) e do referencias.bib.

Entradas (somente leitura): ../../latex/capitulos/*.tex, ../../latex/referencias.bib, templates/template-dissertacao-abnt.docx
Saidas: ../saida/<NOME>-ABNT.docx (o PDF e gerado depois pelo Word, ver atualizar_campos_e_pdf.ps1)
"""
from __future__ import annotations

import re
import sys
import unicodedata
from pathlib import Path

ABNT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ABNT / ".libs"))
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

LATEX = ABNT.parent / "latex"
CAPS = sorted((LATEX / "capitulos").glob("cap*.tex"))
BIBTXT = (LATEX / "referencias.bib").read_text(encoding="utf-8")
OUT = ABNT / "saida"
OUT.mkdir(exist_ok=True)
BS = chr(92)
ACESSO = "10 out. 2026"
NOME = "Raphael_Eiras-Decomposicao_por_aspecto_do_sentimento_em_avaliacoes_do_airbnb_rj-ucp_2026-ABNT"

# ----------------------------------------------------------------------------- bib
ACC = {"'": "\u0301", "`": "\u0300", "^": "\u0302", '"': "\u0308", "~": "\u0303", "c": "\u0327"}


def tex2uni(s: str) -> str:
    s = s.replace(BS + BS, BS)
    s = s.replace("{" + BS + "L}", "Ł").replace("{" + BS + "l}", "ł")
    s = re.sub(r"\{(['`^\"~])([A-Za-z])\}", lambda m: unicodedata.normalize("NFC", m.group(2) + ACC[m.group(1)]), s)
    s = re.sub(r"\\(['`^\"~c])\s*\{?([A-Za-z])\}?", lambda m: unicodedata.normalize("NFC", m.group(2) + ACC[m.group(1)]), s)
    s = re.sub(r"\{\$P_i\$\}", "Pi", s)
    s = s.replace(r"\&", "&")
    s = re.sub(r"\\url\{([^}]*)\}", r"\1", s)
    return re.sub(r"[{}]", "", s).replace("--", "-").strip()


def read_braced(s: str, i: int) -> tuple[str, int]:
    assert s[i] == "{", s[i : i + 20]
    d = 0
    for j in range(i, len(s)):
        if s[j] == "{":
            d += 1
        elif s[j] == "}":
            d -= 1
            if d == 0:
                return s[i + 1 : j], j + 1
    raise ValueError("chave nao fechada")


def parse_bib(txt: str) -> dict:
    out = {}
    for m in re.finditer(r"@(\w+)\{([^,\s]+),", txt):
        if m.group(1).lower() == "string":
            continue
        j = m.end()
        d = 1
        k = j
        while d:
            if txt[k] == "{":
                d += 1
            elif txt[k] == "}":
                d -= 1
            k += 1
        body = txt[j : k - 1]
        f = {}
        for fm in re.finditer(r"(\w+)\s*=\s*", body):
            p = fm.end()
            if p < len(body) and body[p] == "{":
                v, _ = read_braced(body, p)
                f.setdefault(fm.group(1).lower(), v)
        f["_type"] = m.group(1).lower()
        out[m.group(2)] = f
    return out


BIB = parse_bib(BIBTXT)


def split_authors(a: str) -> list[tuple[str, str, bool]]:
    parts, d, cur = [], 0, ""
    toks = re.split(r"(\s+and\s+)", a)
    buf = ""
    for t in toks:
        if re.fullmatch(r"\s+and\s+", t) and buf.count("{") == buf.count("}"):
            parts.append(buf)
            buf = ""
        else:
            buf += t
    parts.append(buf)
    res = []
    for p in parts:
        p = p.strip()
        if p == "others":
            continue
        if p.startswith("{"):
            res.append((tex2uni(p), "", True))
        else:
            fam, _, giv = p.partition(",")
            res.append((tex2uni(fam), tex2uni(giv), False))
    return res


def norm(s: str) -> str:
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()


def info(key: str) -> dict:
    e = BIB[key]
    au = split_authors(e.get("author", ""))
    return {"au": au, "fam": [a[0] for a in au], "year": e["year"], "inst": au[0][2] if au else False}


def cite_p(keys: list[str]) -> str:
    items = []
    for k in keys:
        i = info(k)
        f = list(i["fam"])  # NBR 10520:2023: sobrenome com inicial maiuscula
        txt = "; ".join(f) if len(f) <= 3 else f[0] + " et al."
        items.append((norm(i["fam"][0]), i["year"], f"{txt}, {i['year']}"))
    items.sort()
    return "(" + "; ".join(x[2] for x in items) + ")"


def cite_t(keys: list[str]) -> str:
    out = []
    for k in keys:
        i = info(k)
        f = i["fam"]
        if len(f) == 1:
            n = f[0]
        elif len(f) == 2:
            n = f"{f[0]} e {f[1]}"
        elif len(f) == 3:
            n = f"{f[0]}, {f[1]} e {f[2]}"
        else:
            n = f"{f[0]} et al."
        out.append(f"{n} ({i['year']})")
    return "; ".join(out)


# ----------------------------------------------------------------------------- referencias (NBR 6023)
def fin(t: str) -> str:
    t = t.strip()
    return t if t.endswith((".", "?", "!")) else t + "."


def autores_ref(au) -> str:
    def one(a):
        fam, giv, inst = a
        return fam.upper() if inst else (f"{fam.upper()}, {giv}" if giv else fam.upper())

    if len(au) > 3:
        return one(au[0]) + " et al."
    return "; ".join(one(a) for a in au)


def pags(p: str) -> str:
    return p.replace("--", "-").replace("–", "-")


TITULOS = {"brown1910": "Some experimental results in the correlation of mental abilities"}  # remove o '1' de nota de rodape
OVERRIDES = {"mchugh2012": {"volume": "22", "number": "3"}}  # conferido na 1a pagina do PDF (Biochemia Medica 2012;22(3):276-82)


def ref_runs(key: str) -> list[tuple[str, bool]]:
    e = BIB[key]
    t = e["_type"]
    au = split_authors(e.get("author", ""))
    A = autores_ref(au)
    title = tex2uni(e.get("title", ""))
    title = TITULOS.get(key, title)
    e = {**e, **OVERRIDES.get(key, {})}
    yr = e["year"]
    doi = e.get("doi")
    url = e.get("url")
    R: list[tuple[str, bool]] = []
    if t == "article":
        jn = tex2uni(e["journal"])
        if jn.lower().startswith("arxiv"):
            R = [(f"{fin(A)} {fin(title)} ", False), (jn, True), (f", {yr}. Disponível em: {url}. Acesso em: {ACESSO}.", False)]
        else:
            tail = ""
            if e.get("volume"):
                tail += f", v. {e['volume']}"
            if e.get("number"):
                tail += f", n. {e['number']}"
            if e.get("pages"):
                tail += f", p. {pags(e['pages'])}"
            tail += f", {yr}."
            if doi:
                tail += f" DOI: {doi}."
            elif url:
                tail += f" Disponível em: {url}. Acesso em: {ACESSO}."
            R = [(f"{fin(A)} {fin(title)} ", False), (jn, True), (tail, False)]
    elif t == "inproceedings":
        bt = tex2uni(e["booktitle"])
        pub = tex2uni(e.get("publisher", "")) or "[s. n.]"
        loc = tex2uni(e.get("address", "")).split(",")[0] or "[S. l.]"
        tail = f". [S. l.]: {pub}, {yr}." if loc == "[S. l.]" else f". {loc}: {pub}, {yr}."
        if e.get("pages"):
            tail += f" p. {pags(e['pages'])}."
        if doi:
            tail += f" DOI: {doi}."
        elif url:
            tail += f" Disponível em: {url}. Acesso em: {ACESSO}."
        R = [(f"{fin(A)} {fin(title)} In: ", False), (bt, True), (tail, False)]
    elif t == "mastersthesis":
        sch = tex2uni(e["school"])
        univ_nome, _, univ_unidade = sch.partition(", ")
        R = [(f"{fin(A)} ", False), (title, True), (f". {yr}. Dissertação (Mestrado) – {univ_unidade}, {univ_nome}, {tex2uni(e.get('address', 'Campinas'))}, {yr}.", False)]
        if doi:
            R[-1] = (R[-1][0] + f" DOI: {doi}.", False)
    elif t == "book":
        pub = tex2uni(e.get("publisher", ""))
        loc = tex2uni(e.get("address", "")) or "[S. l.]"
        serie = f" ({tex2uni(e['series'])})" if e.get("series") else ""
        tail = f". {loc}: {pub}, {yr}.{serie}"
        if doi:
            tail += f" DOI: {doi}."
        R = [(f"{fin(A)} ", False), (title, True), (tail, False)]
    elif t == "misc":
        how = tex2uni(e.get("howpublished", ""))
        R = [(f"{fin(A)} ", False), (title, True), (f". {yr}. Disponível em: {how}. Acesso em: {ACESSO}.", False)]
    else:
        R = [(f"{A}. {title}. {yr}.", False)]
    return R


# ----------------------------------------------------------------------------- latex -> runs
LABELS: dict[str, int] = {}


def read_arg(s: str, i: int) -> tuple[str, int]:
    while i < len(s) and s[i] in " \n":
        i += 1
    return read_braced(s, i)


def math_runs(m: str, st: dict) -> list[tuple[str, dict]]:
    m = m.replace(BS + ";", " ").replace(BS + ",", " ").replace("{,}", ",")
    for a, b in [(BS + "cdot", "·"), (BS + "times", "×"), (BS + "rho", "ρ"), (BS + "pm", "±"), (BS + "geq", "≥"), (BS + "leq", "≤"), (BS + "sum", "Σ"), (BS + "min", "min"), (BS + "max", "max")]:
        m = m.replace(a, b)
    m = re.sub(re.escape(BS) + r"text\{([^}]*)\}", r"\1", m)
    runs: list[tuple[str, dict]] = []
    i = 0
    first = True
    while i < len(m):
        c = m[i]
        if c in "^_":
            i += 1
            if i < len(m) and m[i] == "{":
                arg, i = read_braced(m, i)
            else:
                arg, i = m[i], i + 1
            runs.append((arg, {**st, "sub" if c == "_" else "sup": True}))
            continue
        if m.startswith(("min", "max"), i):
            runs.append((m[i : i + 3], dict(st)))
            i += 3
            first = False
            continue
        if c == "-" and (first or m[i - 1] in " (["):
            runs.append(("−", dict(st)))
        elif c.isalpha():
            nxt = m[i + 1] if i + 1 < len(m) else ""
            runs.append((c, {**st, "i": True} if not nxt.isalpha() else dict(st)))
        else:
            runs.append((c, dict(st)))
        first = False
        i += 1
    return runs


def parse_inline(s: str, st: dict | None = None) -> list[tuple[str, dict]]:
    st = dict(st or {})
    out: list[tuple[str, dict]] = []
    buf = ""
    i = 0

    def flush():
        nonlocal buf
        if buf:
            out.append((buf, dict(st)))
            buf = ""

    while i < len(s):
        c = s[i]
        if c == "$":
            j = s.index("$", i + 1)
            flush()
            out.extend(math_runs(s[i + 1 : j], st))
            i = j + 1
            continue
        if c == BS:
            m = re.match(re.escape(BS) + r"([a-zA-Z]+)", s[i:])
            if m:
                name = m.group(1)
                i += len(m.group(0))
                if name in ("textit", "emph"):
                    arg, i = read_arg(s, i)
                    flush()
                    out.extend(parse_inline(arg, {**st, "i": True}))
                elif name == "textbf":
                    arg, i = read_arg(s, i)
                    flush()
                    out.extend(parse_inline(arg, {**st, "b": True}))
                elif name == "texttt":
                    arg, i = read_arg(s, i)
                    flush()
                    out.extend(parse_inline(arg, {**st, "mono": True}))
                elif name == "path":
                    arg, i = read_arg(s, i)
                    flush()
                    out.append((arg, {**st, "mono": True}))
                elif name in ("citep", "citet"):
                    arg, i = read_arg(s, i)
                    keys = [k.strip() for k in arg.split(",")]
                    flush()
                    out.append(((cite_p if name == "citep" else cite_t)(keys), dict(st)))
                elif name == "ref":
                    arg, i = read_arg(s, i)
                    buf += str(LABELS[arg])
                elif name == "label":
                    _, i = read_arg(s, i)
                elif name == "ajusta":
                    arg, i = read_arg(s, i)
                    flush()
                    out.extend(parse_inline(arg, st))
                elif name in ("allowbreak", "small", "scriptsize", "centering", "hline", "normalsize"):
                    while i < len(s) and s[i] == " ":
                        i += 1
                continue
            nxt = s[i + 1] if i + 1 < len(s) else ""
            if nxt in "%&_#{}$":
                buf += nxt
            elif nxt in ";,":
                buf += " "
            elif nxt == BS:
                buf += "\n"
            i += 2
            continue
        if c == "~":
            buf += "\u00a0"
        elif c in "{}":
            pass
        elif s.startswith("--", i):
            buf += "–"
            i += 2
            continue
        elif c == "\n":
            buf += " "
        else:
            buf += c
        i += 1
    flush()
    return out


def add_runs(par, runs, size=None, base_bold=False):
    for text, st in runs:
        parts = text.split("\n")
        for n, part in enumerate(parts):
            if n:
                par.add_run().add_break()
            if not part:
                continue
            r = par.add_run(part)
            if st.get("i"):
                r.italic = True
            if st.get("b") or base_bold:
                r.bold = True
            if st.get("mono"):
                r.font.name = "Courier New"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
                r.font.size = Pt(10 if size is None else min(size, 10))
            elif size:
                r.font.size = Pt(size)
            if st.get("sup"):
                r.font.superscript = True
            if st.get("sub"):
                r.font.subscript = True


# ----------------------------------------------------------------------------- docx helpers
def set_par_text(p, text: str):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)


def remove(p):
    p._element.getparent().remove(p._element)


def field_runs(par, instr: str, cached: str, size=None):
    def fc(t):
        r = par.add_run()
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), t)
        r._r.append(e)
        return r

    fc("begin")
    r = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fc("separate")
    rr = par.add_run(cached)
    if size:
        rr.font.size = Pt(size)
    fc("end")


def cell_borders(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side, sz in (("top", top), ("bottom", bottom)):
        if sz:
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "000000")
            b.append(e)
    tcPr.append(b)


def forcar_times(doc):
    """Remove as fontes de tema (Calibri/Cambria) de todos os estilos e fixa Times New Roman."""
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for rf in doc.styles.element.iter("{%s}rFonts" % W):
        for att in list(rf.attrib):
            if att.lower().endswith("theme"):
                del rf.attrib[att]
        for att in ("ascii", "hAnsi", "cs", "eastAsia"):
            rf.set("{%s}%s" % (W, att), "Times New Roman")


def page_break_par(anchor):
    p = anchor.insert_paragraph_before("")
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


# ----------------------------------------------------------------------------- partes do texto
def strip_comments(t: str) -> str:
    return re.sub(r"(?m)^\s*%.*$", "", t)


def table_blocks(text: str):
    """Quebra o tex de um capitulo em blocos: ('h1', titulo) ('h2', titulo) ('p', texto) ('table', dict)."""
    text = strip_comments(text)
    i = 0
    blocks = []
    paras: list[str] = []

    def flush():
        if paras:
            blocks.append(("p", " ".join(paras)))
            paras.clear()

    lines = text.split("\n")
    n = 0
    while n < len(lines):
        ln = lines[n].strip()
        if not ln:
            flush()
        elif ln.startswith(BS + "section{"):
            flush()
            blocks.append(("h1", read_braced(ln, ln.index("{"))[0]))
        elif ln.startswith(BS + "subsection{"):
            flush()
            blocks.append(("h2", read_braced(ln, ln.index("{"))[0]))
        elif ln == BS + "appendix":
            flush()
            blocks.append(("appendix", None))
        elif ln.startswith(BS + "begin{figure}"):
            flush()
            buf = []
            while not lines[n].strip().startswith(BS + "end{figure}"):
                buf.append(lines[n])
                n += 1
            body = "\n".join(buf)
            ci = body.index(BS + "caption{") + len(BS + "caption")
            cap, _ = read_braced(body, ci)
            lab = re.search(re.escape(BS) + r"label\{(.*?)\}", body).group(1)
            ig = re.search(re.escape(BS) + r"includegraphics\[(.*?)\]\{(.*?)\}", body)
            wm = re.search(r"width=([0-9.]+)", ig.group(1))
            blocks.append(("figure", {"caption": cap, "label": lab, "arquivo": ig.group(2), "largura": float(wm.group(1)) if wm else 1.0}))
        elif ln.startswith(BS + "begin{table}"):
            flush()
            buf = []
            while not lines[n].strip().startswith(BS + "end{table}"):
                buf.append(lines[n])
                n += 1
            body = "\n".join(buf)
            ci = body.index(BS + "caption{") + len(BS + "caption")
            cap, _ = read_braced(body, ci)
            lab = re.search(re.escape(BS) + r"label\{(.*?)\}", body).group(1)
            ti = body.index(BS + "begin{tabular}") + len(BS + "begin{tabular}")
            spec, after = read_braced(body, ti)
            rows = body[after : body.index(BS + "end{tabular}")]
            blocks.append(("table", {"caption": cap, "label": lab, "spec": spec, "rows": rows}))
        else:
            paras.append(ln)
        n += 1
    flush()
    return blocks


def spec_align(spec: str, ncols: int) -> list[str]:
    toks = re.findall(r"p\{[^}]*\}|[lcr]", spec.replace("|", ""))
    al = []
    for t in toks:
        al.append("l" if t.startswith("p") or t == "l" else t)
    return (al + ["r"] * ncols)[:ncols]


def build_table(doc, anchor, tb: dict, num: int, larg_cm: float = 16.0):
    # legenda (acima)
    cp = anchor.insert_paragraph_before("", style="Caption")
    cp.paragraph_format.keep_with_next = True
    cp.paragraph_format.space_before = Pt(6)
    cp.paragraph_format.space_after = Pt(4)
    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = cp.add_run("Tabela ")
    r.font.size = Pt(10)
    r.bold = False
    field_runs(cp, "SEQ Tabela \\* ARABIC", str(num), 10)
    add_runs(cp, [(" – ", {})] + parse_inline(tb["caption"]), size=10)
    for r in cp.runs:
        r.bold = False
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0, 0, 0)
    rows = []
    for raw in tb["rows"].replace(BS + "hline", "").split(BS + BS):
        raw = raw.strip()
        if raw:
            rows.append([c.strip() for c in raw.split("&")])
    ncols = max(len(r) for r in rows)
    al = spec_align(tb["spec"], ncols)
    table = doc.add_table(rows=len(rows), cols=ncols)
    anchor._p.addprevious(table._tbl)
    table.alignment = 1
    fs = 10 if ncols <= 4 else (8.5 if ncols <= 6 else 8)
    w = [max(len(re.sub(r"[$" + re.escape(BS) + "{}]", "", c[ci])) if ci < len(c) else 0 for c in rows) for ci in range(ncols)]
    w = [min(max(x, 6), 42) + 3 for x in w]
    tot = sum(w)
    widths = [Cm(larg_cm * x / tot) for x in w]
    table.autofit = False
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.width = widths[ci]
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.first_line_indent = Cm(0)
            par.paragraph_format.keep_with_next = True
            par.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}[al[ci]]
            txt = row[ci] if ci < len(row) else ""
            add_runs(par, parse_inline(txt), size=fs, base_bold=(ri == 0))
            if ri == 0:
                cell_borders(cell, top=8, bottom=4)
            if ri == len(rows) - 1:
                cell_borders(cell, bottom=8)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    fp = anchor.insert_paragraph_before("")
    fp.paragraph_format.space_before = Pt(3)
    fp.paragraph_format.space_after = Pt(12)
    fp.paragraph_format.line_spacing = 1.0
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.paragraph_format.first_line_indent = Cm(0)
    rr = fp.add_run("Fonte: Elaborada pelo autor.")
    rr.font.size = Pt(10)
    return table


def main():
    # rotulos das tabelas, na ordem de aparecimento
    k = 0
    for c in CAPS:
        for kind, val in table_blocks(c.read_text(encoding="utf-8")):
            if kind == "table":
                k += 1
                LABELS[val["label"]] = k
    resumo = parse_inline(strip_comments((LATEX / "capitulos" / "resumo.tex").read_text(encoding="utf-8")).strip().replace("\n", " "))
    titulo_lines = (
        "O que as avaliações revelam sobre a nota",
        "decomposição por aspecto do sentimento em avaliações de hóspedes do Airbnb no Rio de Janeiro",
    )
    autor = "Raphael Rodrigues Eiras"
    orient = "Prof. Dr. Rodolfo Nicolay"
    prog = "Programa de Mestrado Profissional em Gestão de Sistemas de Engenharia"
    univ = "Universidade Católica de Petrópolis"
    natureza = f"Dissertação apresentada ao {prog} da {univ}, como requisito parcial para a obtenção do título de Mestre."
    palavras = "análise de sentimento por aspecto; avaliações online; Airbnb; modelagem de tópicos; diagnóstico de anúncios"
    abstract = (
        "Ratings on Airbnb are concentrated at the top of the scale and say little about the reasons for guest satisfaction. "
        "This study uses the text of 995,072 reviews of listings in Rio de Janeiro (Inside Airbnb) to decompose the rating by aspect and by specific complaint. "
        "Sentence-level sentiment is measured by two multilingual models, and the choice of the main model follows a rule fixed before the second model was run. "
        "Aspects come from topics obtained with BERTopic and are checked with a second method (LDA) and with a second annotator. "
        "The measurement is validated with the official subscores and with labeled sentences weighted by the population. "
        "Only cleanliness, location and value meet the adopted confidence criterion. "
        "Amenities and cleanliness concentrate most of the explained variation of the overall rating. "
        "Complaints about facilities, cleanliness and host response are the most associated with lower ratings, with a similar order in both models. "
        "A synthetic index built from the text adds no explanatory power, beyond the official rating, over indices of review volume. "
        "The results are associations, and the proposed application is a listing-level diagnosis compared with the neighborhood and the type of accommodation."
    )
    keywords = "aspect-based sentiment analysis; online reviews; Airbnb; topic modeling; listing diagnosis"

    doc = Document(str(ABNT / "templates" / "template-dissertacao-abnt.docx"))
    P = list(doc.paragraphs)

    def sub_title(p):
        p.clear()
        r = p.add_run(titulo_lines[0].upper())
        r2 = p.add_run(": " + titulo_lines[1])
        r2.font.bold = False
        r2.font.all_caps = False

    # --- capa
    set_par_text(P[0], univ.upper())
    remove(P[1])  # unidade/faculdade: nao confirmada no manual
    set_par_text(P[2], prog.upper())
    set_par_text(P[3], autor.upper())
    sub_title(P[4])
    set_par_text(P[5], "Petrópolis")
    set_par_text(P[6], "2026")
    # --- folha de rosto
    set_par_text(P[8], autor.upper())
    sub_title(P[9])
    set_par_text(P[10], natureza)
    set_par_text(P[11], f"Orientador: {orient}")
    remove(P[12])  # sem coorientador
    set_par_text(P[13], "Petrópolis")
    P[13].paragraph_format.space_before = Pt(200)  # local e ano no pe da folha de rosto
    set_par_text(P[14], "2026")
    # --- aprovacao
    set_par_text(P[19], autor.upper())
    sub_title(P[20])
    set_par_text(P[21], natureza)
    set_par_text(P[22], "Aprovada em: [dia] de [mês] de [ano].")
    set_par_text(P[25], f"{orient} – {univ} (orientador)")
    # --- agradecimentos (opcional): removidos
    for idx in (31, 32, 33):
        remove(P[idx])
    # --- resumo
    p = P[35]
    p.clear()
    p.paragraph_format.first_line_indent = Cm(0)
    add_runs(p, resumo)
    set_par_text(P[36], f"Palavras-chave: {palavras}.")
    P[36].paragraph_format.first_line_indent = Cm(0)
    # --- abstract
    P[39].clear()
    P[39].paragraph_format.first_line_indent = Cm(0)
    P[39].add_run(abstract)
    set_par_text(P[40], f"Keywords: {keywords}.")
    P[40].paragraph_format.first_line_indent = Cm(0)
    # --- lista de tabelas (antes do sumario)
    anchor = P[41]  # quebra de pagina antes do sumario
    tit = anchor.insert_paragraph_before("LISTA DE TABELAS", style="Título não numerado")
    lt = anchor.insert_paragraph_before("")
    lt.paragraph_format.first_line_indent = Cm(0)
    lt.paragraph_format.line_spacing = 1.0
    caps_titles = []
    for c in CAPS:
        for kind, val in table_blocks(c.read_text(encoding="utf-8")):
            if kind == "table":
                caps_titles.append("".join(t for t, _ in parse_inline(val["caption"])))
    cached = "\n".join(f"Tabela {n} – {t}" for n, t in enumerate(caps_titles, 1))
    field_runs(lt, 'TOC \\h \\z \\c "Tabela"', cached, 11)
    # a lista tem de ficar numa pagina propria, antes da quebra do sumario
    page_break_par(tit)
    # --- corpo
    anchor = P[56]  # REFERENCIAS
    for idx in range(45, 56):
        remove(P[idx])
    h1 = h2 = 0
    tnum = 0
    for c in CAPS:
        for kind, val in table_blocks(c.read_text(encoding="utf-8")):
            if kind == "h1":
                h1 += 1
                h2 = 0
                anchor.insert_paragraph_before(f"{h1} {val.upper()}", style="Heading 1")
            elif kind == "h2":
                h2 += 1
                anchor.insert_paragraph_before(f"{h1}.{h2} {val.upper()}", style="Heading 2")
            elif kind == "p":
                par = anchor.insert_paragraph_before("", style="Corpo do texto")
                add_runs(par, parse_inline(val))
            elif kind == "table":
                tnum += 1
                build_table(doc, anchor, val, tnum)
    # --- referencias
    keys = sorted(BIB, key=lambda k: (norm(info(k)["fam"][0]), info(k)["year"], norm(BIB[k].get("title", ""))))
    ref_anchor = P[59]  # titulo do apendice (sera removido); entradas entram antes dele
    for k in keys:
        par = ref_anchor.insert_paragraph_before("", style="Referência")
        for text, bold in ref_runs(k):
            r = par.add_run(text)
            r.bold = bold
    remove(P[57])
    remove(P[58])
    remove(P[59])
    remove(P[60])
    forcar_times(doc)
    out = OUT / (NOME + ".docx")
    doc.save(str(out))
    print("salvo:", out, "| tabelas:", tnum, "| referencias:", len(keys))


if __name__ == "__main__":
    main()
