"""Fase 4 — preparação do segundo anotador e cálculo do kappa entre anotadores.

`gerar`: sorteia 100 das 297 frases já rotuladas (mesma estratificação: aspecto previsto × polaridade prevista), embaralha, e grava
  - data/validacao_segundo_anotador.xlsx  -> ENVIAR ao segundo anotador: só o texto das frases (sem rótulos do primeiro anotador,
    sem previsões de modelo, sem idioma), com instruções e definições de cada aspecto;
  - data/validacao_segundo_anotador_CHAVE.xlsx -> NÃO ENVIAR: liga o id cego ao id da amostra original e traz os rótulos do primeiro anotador.
`kappa`: lê a planilha preenchida e calcula a concordância entre anotadores (sentimento e aspecto). Nunca sobrescreve planilhas existentes.
"""
from __future__ import annotations

import logging
from pathlib import Path

import numpy as np
import pandas as pd
from sklearn.metrics import cohen_kappa_score, confusion_matrix

from config import GENERAL_BUCKET

log = logging.getLogger("pipeline")

DEFINICOES = {
    "limpeza": "Limpeza e higiene do imóvel: sujeira, poeira, cheiro, mofo, pragas, roupa de cama e toalhas limpas, estado de limpeza na chegada.",
    "localizacao": "Onde o imóvel fica: proximidade de praia, metrô, restaurantes, pontos turísticos; facilidade de chegar a pé ou de transporte; qualidade do bairro como local.",
    "checkin": "Chegada e saída: processo de check-in e check-out, horários, flexibilidade, chaves, fechadura eletrônica, portaria na chegada, acesso ao imóvel.",
    "comunicacao": "Comunicação e atendimento do anfitrião ou de quem o representa: rapidez de resposta, simpatia, ajuda, dicas, hospitalidade, equipe do prédio.",
    "custo_beneficio": "Relação entre o preço pago e o que se recebeu: 'vale o preço', 'barato', 'caro', custo-benefício.",
    "ruido_conforto": "Ruído e conforto de dormir: barulho da rua ou de vizinhos, silêncio, luz e claridade, sono. Não inclui o equipamento em si (ver comodidades).",
    "seguranca": "Segurança: sensação de segurança, assaltos, portaria 24 horas como segurança, segurança do prédio e do entorno.",
    "precisao": "Fidelidade ao anúncio: o imóvel é como nas fotos e na descrição; diferenças entre o anunciado e o real.",
    "comodidades": "Instalações e equipamentos: cama e colchão, banheiro e chuveiro, ar-condicionado, cozinha e utensílios, wifi, TV, piscina, varanda, vista, máquina de lavar, decoração, estado e manutenção do imóvel.",
    "satisfacao_geral": "Elogio ou queixa GENÉRICO sobre a experiência, sem tratar de um aspecto específico: 'recomendo', 'adorei', 'voltarei', 'estadia incrível', 'decepcionante'.",
    "geral": "Frase que não trata do imóvel nem da estadia (informação solta, outro assunto, agradecimento sem conteúdo) ou que não dá para classificar. Use também no lugar de 'nenhum'.",
}
SENT_DEF = ("negativo: a frase, isolada, expressa insatisfação ou crítica sobre o imóvel/estadia. neutro: informa ou descreve sem avaliação clara, ou mistura sem predomínio. "
            "positivo: expressa satisfação ou elogio.")


def _estratificar(H: pd.DataFrame, n: int, seed: int) -> pd.DataFrame:
    """100 de 297 com a mesma estratificação (aspecto previsto × polaridade prevista): ~3 por célula, completando com sorteio."""
    rng = np.random.default_rng(seed)
    pol = pd.cut(H["estrelas_modelo"], [0, 2.5, 3.5, 5.01], labels=["neg", "neu", "pos"])
    H = H.assign(pol=pol)
    cells = list(H.groupby(["aspecto_modelo", "pol"], observed=True).groups.items())
    por_celula = max(1, n // len(cells))
    escolhidos = []
    for _, idx in cells:
        idx = list(idx)
        rng.shuffle(idx)
        escolhidos += idx[:por_celula]
    resto = [i for i in H.index if i not in set(escolhidos)]
    rng.shuffle(resto)
    escolhidos += resto[: max(0, n - len(escolhidos))]
    return H.loc[escolhidos[:n]]


def gerar(cfg, n: int = 100) -> tuple[Path, Path]:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    src = Path(cfg.MANUAL_VALIDATION_XLSX)
    H = pd.read_excel(src, sheet_name="Amostra").merge(pd.read_excel(src, sheet_name="Gabarito_modelo"), on="id_amostra")
    S = _estratificar(H, n, cfg.SEED + 7).sample(frac=1, random_state=cfg.SEED + 8).reset_index(drop=True)
    S.insert(0, "id_cego", np.arange(1, len(S) + 1))
    out = Path(cfg.MANUAL_VALIDATION_XLSX).with_name("validacao_segundo_anotador.xlsx")
    key = out.with_name("validacao_segundo_anotador_CHAVE.xlsx")
    if out.exists() or key.exists():
        raise SystemExit(f"{out.name} ou {key.name} já existe; não sobrescrevo (pode conter rótulos reais).")

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Instrucoes"
    linhas = ["Tarefa: rotule cada frase da aba 'Amostra' de forma independente (sem consultar ninguém e sem ver classificações automáticas).",
              "Contexto: as frases são sentenças de avaliações de hóspedes de imóveis do Airbnb no Rio de Janeiro (português, inglês, espanhol, francês etc.).",
              "", "Coluna sentimento_humano (negativo / neutro / positivo): " + SENT_DEF,
              "Coluna aspecto_humano: o aspecto PRINCIPAL da frase (escolha na lista). Se a frase tocar em mais de um, escolha o mais importante para o hóspede.",
              "", "Definições dos aspectos:"] + [f"  - {k}: {v}" for k, v in DEFINICOES.items()] + \
             ["", "Frases curtas e ambíguas existem; escolha a melhor opção e use a coluna 'observacao' se quiser comentar.",
              "Salve o arquivo com o mesmo nome e devolva ao pesquisador."]
    for i, t in enumerate(linhas, 1):
        ws0.cell(i, 1, t)
    ws0.column_dimensions["A"].width = 160
    ws = wb.create_sheet("Amostra")
    ws.append(["id_cego", "texto", "sentimento_humano", "aspecto_humano", "observacao"])
    for r in S.itertuples():
        ws.append([r.id_cego, r.texto, None, None, None])
    m = len(S) + 1
    dv1 = DataValidation(type="list", formula1='"negativo,neutro,positivo"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(list(DEFINICOES)) + '"', allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"C2:C{m}")
    dv2.add(f"D2:D{m}")
    ws.column_dimensions["B"].width = 110
    ws.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    S[["id_cego", "id_amostra", "sentimento_humano", "aspecto_humano", "aspecto_modelo", "estrelas_modelo"]].rename(
        columns={"sentimento_humano": "sentimento_anotador1", "aspecto_humano": "aspecto_anotador1"}).to_excel(key, index=False)
    log.info("Planilha do segundo anotador: %s (%d frases); chave: %s (NÃO ENVIAR)", out, len(S), key)
    return out, key


def kappa(cfg, preenchida: str | Path, chave: str | Path | None = None, boot: int = 2000) -> dict[str, pd.DataFrame]:
    preenchida = Path(preenchida)
    chave = Path(chave) if chave else preenchida.with_name("validacao_segundo_anotador_CHAVE.xlsx")
    A2 = pd.read_excel(preenchida, sheet_name="Amostra")
    K = pd.read_excel(chave)
    D = K.merge(A2[["id_cego", "sentimento_humano", "aspecto_humano"]].rename(columns={"sentimento_humano": "sentimento_anotador2", "aspecto_humano": "aspecto_anotador2"}), on="id_cego")
    D = D[D["sentimento_anotador2"].notna() & D["aspecto_anotador2"].notna()].copy()
    if len(D) < 20:
        raise SystemExit(f"Apenas {len(D)} frases rotuladas pelo segundo anotador; o kappa não é calculado com menos de 20.")
    for c in ("aspecto_anotador1", "aspecto_anotador2"):
        D[c] = D[c].replace({"nenhum": GENERAL_BUCKET})  # 'nenhum' e 'geral' são o mesmo grupo no estudo
    num = {"negativo": -1, "neutro": 0, "positivo": 1}
    s1, s2 = D["sentimento_anotador1"].map(num).to_numpy(), D["sentimento_anotador2"].map(num).to_numpy()
    a1, a2 = D["aspecto_anotador1"].to_numpy(), D["aspecto_anotador2"].to_numpy()
    rng = np.random.default_rng(cfg.SEED)

    def boot_ci(f):
        v = [f(i) for i in (rng.integers(0, len(D), len(D)) for _ in range(boot))]
        v = [x for x in v if np.isfinite(x)]
        return float(np.percentile(v, 2.5)), float(np.percentile(v, 97.5))

    ks = lambda i: cohen_kappa_score(s1[i], s2[i])
    kw = lambda i: cohen_kappa_score(s1[i], s2[i], weights="linear")
    ka = lambda i: cohen_kappa_score(a1[i], a2[i])
    todos = np.arange(len(D))
    R = pd.DataFrame([
        {"medida": "n frases rotuladas pelos dois", "valor": len(D), "ic_inf": np.nan, "ic_sup": np.nan},
        {"medida": "sentimento: concordância simples", "valor": float((s1 == s2).mean()), "ic_inf": np.nan, "ic_sup": np.nan},
        {"medida": "sentimento: kappa de Cohen", "valor": float(ks(todos)), **dict(zip(["ic_inf", "ic_sup"], boot_ci(ks)))},
        {"medida": "sentimento: kappa ponderado linear", "valor": float(kw(todos)), **dict(zip(["ic_inf", "ic_sup"], boot_ci(kw)))},
        {"medida": "aspecto: concordância simples", "valor": float((a1 == a2).mean()), "ic_inf": np.nan, "ic_sup": np.nan},
        {"medida": "aspecto: kappa de Cohen", "valor": float(ka(todos)), **dict(zip(["ic_inf", "ic_sup"], boot_ci(ka)))},
    ])
    por_aspecto = D.assign(igual=(a1 == a2)).groupby("aspecto_anotador1")["igual"].agg(["size", "mean"]).rename(columns={"size": "n", "mean": "concordancia"}).reset_index()
    return {"resumo": R, "por_aspecto": por_aspecto, "conf_sentimento": pd.crosstab(D["sentimento_anotador1"], D["sentimento_anotador2"]).reset_index(),
            "conf_aspecto": pd.crosstab(D["aspecto_anotador1"], D["aspecto_anotador2"]).reset_index()}


def salvar_kappa(cfg, store, preenchida: str | Path = "data/validacao_segundo_anotador.xlsx") -> dict[str, pd.DataFrame]:
    """Calcula a concordância entre anotadores e guarda as tabelas no cache (prefixo cmp_, usadas na aba Segundo_Anotador)."""
    r = kappa(cfg, preenchida)
    for k, v in r.items():
        store.save(f"cmp_segundo_anotador_{k}", v)
    return r
