"""Etapa 14 — Validade de construto ampliada e amostra de validação manual.

Pergunta de pesquisa 1: o sentimento por aspecto é adequado às notas e subnotas oficiais? Análises (nível do anúncio):
  (a) concordância por faixa de volume (≤ 30 avaliações: o sentimento usa TODAS; > 30: só as 30 mais recentes);
  (b) matriz convergente/discriminante (cada aspecto × todas as subnotas) e correlação parcial controlando a nota global
      (as subnotas são muito correlacionadas entre si, "efeito halo");
  (c) confiabilidade split-half do sentimento por aspecto e correlação corrigida pela atenuação (só do lado do sentimento:
      a confiabilidade das subnotas oficiais é desconhecida, então a correção é um limite inferior da correlação verdadeira);
  (d) concordância com subnotas saturadas: AUC (subnota = 5 vs < 5) e Kendall tau-b (empates tratados);
  (e) recortes por tipo de acomodação e por idioma das sentenças (validade entre idiomas);
  (f) amostra de validação manual: gera a planilha para rotulagem e, depois de preenchida, calcula acurácia/kappa.
Tudo é associação; nenhuma afirmação causal.
"""
from __future__ import annotations

import logging
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats

from config import GENERAL_BUCKET
from .utils import Store, stage

log = logging.getLogger("pipeline")


def _sp(x, y) -> tuple[float, float, int]:
    d = pd.concat([x, y], axis=1).dropna()
    if len(d) < 30 or d.iloc[:, 0].nunique() < 2 or d.iloc[:, 1].nunique() < 2:
        return np.nan, np.nan, len(d)
    return float(stats.spearmanr(d.iloc[:, 0], d.iloc[:, 1])[0]), float(stats.pearsonr(d.iloc[:, 0], d.iloc[:, 1])[0]), len(d)


def _partial_spearman(x, y, z) -> tuple[float, int]:
    d = pd.concat([x, y, z], axis=1).dropna()
    if len(d) < 30:
        return np.nan, len(d)
    r = d.rank().corr()  # Pearson dos postos = Spearman
    rxy, rxz, ryz = r.iloc[0, 1], r.iloc[0, 2], r.iloc[1, 2]
    return float((rxy - rxz * ryz) / np.sqrt((1 - rxz ** 2) * (1 - ryz ** 2))), len(d)


def por_volume(cfg, D: pd.DataFrame) -> pd.DataFrame:
    bins = [(2, 5), (6, 10), (11, 20), (21, 30), (31, 60), (61, 10 ** 6)]
    rows = []
    for a, sub in cfg.ASPECT_TO_SUBSCORE.items():
        cuts = [(f"v {lo}-{hi if hi < 10 ** 6 else '+'}", (D["number_of_reviews"] >= lo) & (D["number_of_reviews"] <= hi)) for lo, hi in bins]
        cuts += [("v ≤ 30 (sentimento usa todas as avaliações)", D["number_of_reviews"] <= 30),
                 ("v > 30 (só as 30 mais recentes)", D["number_of_reviews"] > 30), ("todos", D["number_of_reviews"] >= 0)]
        for nome, m in cuts:
            rs, rp, n = _sp(D.loc[m, a], D.loc[m, sub])
            rows.append({"aspecto": a, "subnota": sub, "faixa": nome, "n": n, "spearman": rs, "pearson": rp})
    return pd.DataFrame(rows)


def matriz(cfg, D: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    subs = list(dict.fromkeys(cfg.ASPECT_TO_SUBSCORE.values()))
    rows = []
    for a in cfg.ASPECT_SEEDS:
        for s in subs:
            rs, _, n = _sp(D[a], D[s])
            rows.append({"aspecto": a, "subnota": s, "n": n, "spearman": rs})
    M = pd.DataFrame(rows)
    res = []
    for a, own in cfg.ASPECT_TO_SUBSCORE.items():
        r = M[M["aspecto"] == a].set_index("subnota")["spearman"]
        part, n = (np.nan, 0) if own == "review_scores_rating" else _partial_spearman(D[a], D[own], D["review_scores_rating"])
        res.append({"aspecto": a, "subnota_correspondente": own, "spearman_correspondente": r[own],
                    "media_spearman_demais_subnotas": r.drop(own).mean(), "posto_da_correspondente(1=maior)": int(r.rank(ascending=False)[own]),
                    "n_subnotas": len(r), "spearman_parcial_controlando_nota_global": part, "n_parcial": n})
    return M, pd.DataFrame(res)


def auc_kendall(cfg, D: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for a, sub in cfg.ASPECT_TO_SUBSCORE.items():
        d = D[[a, sub]].dropna()
        hi, lo = d.loc[d[sub] >= 5.0, a], d.loc[d[sub] < 5.0, a]
        auc = float(stats.mannwhitneyu(hi, lo, alternative="two-sided").statistic / (len(hi) * len(lo))) if len(hi) and len(lo) else np.nan
        rows.append({"aspecto": a, "subnota": sub, "n": len(d), "n_subnota_igual_5": len(hi), "n_subnota_menor_5": len(lo),
                     "AUC(sentimento maior quando subnota=5)": auc,
                     "dif_media_sentimento(5 - <5)": float(hi.mean() - lo.mean()) if len(hi) and len(lo) else np.nan,
                     "kendall_tau_b": float(stats.kendalltau(d[a], d[sub], variant="b")[0])})
    return pd.DataFrame(rows)


def confiabilidade(cfg, store: Store, D: pd.DataFrame) -> pd.DataFrame:
    """Split-half por AVALIAÇÃO (sentenças da mesma avaliação são correlacionadas), Spearman-Brown, correlação corrigida."""
    S = store.load("sentencas_aspecto")
    S = S[S["aspecto"] != GENERAL_BUCKET]
    rng = np.random.default_rng(cfg.SEED)
    rev = S[["listing_id", "review_id"]].drop_duplicates()
    rev["half"] = rng.integers(0, 2, len(rev))
    S = S.merge(rev, on=["listing_id", "review_id"])
    g = S.groupby(["listing_id", "aspecto", "half"])["stars"].agg(["mean", "count"]).unstack("half")
    rows = []
    for a, sub in cfg.ASPECT_TO_SUBSCORE.items():
        try:
            x = g.xs(a, level="aspecto")
        except KeyError:
            continue
        x = x[(x[("count", 0)] >= cfg.MIN_SENT_PER_ASPECT) & (x[("count", 1)] >= cfg.MIN_SENT_PER_ASPECT)]
        if len(x) < 50:
            continue
        r_hh = float(np.corrcoef(x[("mean", 0)], x[("mean", 1)])[0, 1])
        rel = 2 * r_hh / (1 + r_hh)  # Spearman-Brown para o comprimento total
        sub_d = D.set_index("id").loc[x.index.intersection(D["id"]), [a, sub]].dropna()
        rs, _, n = _sp(sub_d[a], sub_d[sub])
        rows.append({"aspecto": a, "subnota": sub, "n_anuncios_com_split": len(x), "r_split_half": r_hh, "confiabilidade_SB": rel,
                     "n_correlacao": n, "spearman_observado": rs,
                     "spearman_corrigido_atenuacao_lado_sentimento": rs / np.sqrt(rel) if rel > 0 else np.nan})
    return pd.DataFrame(rows)


def recortes(cfg, store: Store, D: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for rt in ("Entire home/apt", "Private room"):
        d = D[D["room_type"] == rt]
        for a, sub in cfg.ASPECT_TO_SUBSCORE.items():
            rs, rp, n = _sp(d[a], d[sub])
            rows.append({"recorte": f"tipo: {rt}", "aspecto": a, "subnota": sub, "n": n, "spearman": rs, "pearson": rp})
    S = store.load("sentencas_aspecto")
    S = S[S["aspecto"] != GENERAL_BUCKET]
    L = D.set_index("id")
    for lang in ("pt", "en", "es"):
        g = S[S["lang"] == lang].groupby(["listing_id", "aspecto"])["stars"].agg(["mean", "count"]).reset_index()
        g = g[g["count"] >= cfg.MIN_SENT_PER_ASPECT]
        for a, sub in cfg.ASPECT_TO_SUBSCORE.items():
            x = g[g["aspecto"] == a].set_index("listing_id")["mean"]
            rs, rp, n = _sp(x.rename("x"), L[sub].reindex(x.index).rename("y"))
            rows.append({"recorte": f"idioma das sentenças: {lang}", "aspecto": a, "subnota": sub, "n": n, "spearman": rs, "pearson": rp})
    return pd.DataFrame(rows)


# ------------------------------------------------------------------ amostra manual
def gerar_amostra_manual(cfg, store: Store, path: Path) -> None:
    """Amostra estratificada (aspecto × polaridade prevista). NUNCA sobrescreve uma planilha existente (preserva rótulos)."""
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    S = store.load("sentencas_aspecto")
    T = store.load("texts")[["text_id", "text", "n_tok"]]
    Sn = store.load("sentiment")[["text_id", "p_neg", "p_neu", "p_pos"]]
    S = S.merge(T, on="text_id").merge(Sn, on="text_id")
    S = S[S["n_tok"] >= cfg.MIN_SENT_TOKENS].drop_duplicates("text_id")
    groups = list(cfg.ASPECT_SEEDS) + [GENERAL_BUCKET]
    per = max(3, cfg.MANUAL_SAMPLE_N // len(groups))
    S["pol"] = pd.cut(S["stars"], [0, 2.5, 3.5, 5.01], labels=["neg", "neu", "pos"])
    parts = []
    for gname in groups:
        d = S[S["aspecto"] == gname]
        for pol in ("neg", "neu", "pos"):
            dd = d[d["pol"] == pol]
            parts.append(dd.sample(min(len(dd), per // 3), random_state=cfg.SEED))
    A = pd.concat(parts).sample(frac=1, random_state=cfg.SEED).reset_index(drop=True)
    A.insert(0, "id_amostra", np.arange(1, len(A) + 1))

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Instrucoes"
    for i, line in enumerate([
        "Rotule SEM olhar a aba Gabarito_modelo (evita ancoragem).",
        "sentimento_humano: negativo / neutro / positivo (a frase, isolada, expressa que sentimento sobre o imóvel/estadia?).",
        "aspecto_humano: o aspecto principal da frase (lista). Elogio/queixa genérico sobre a experiência (ex.: 'recomendo', "
        "'adorei') = satisfacao_geral. Use 'geral' ou 'nenhum' para frases que não tratam do imóvel/estadia (o estudo agrupa os dois).",
        "Salve o arquivo com o mesmo nome. Rode: python run_pipeline.py --stages 14  (calcula acurácia, kappa e matrizes de confusão).",
    ], 1):
        ws0.cell(i, 1, line)
    ws = wb.create_sheet("Amostra")
    ws.append(["id_amostra", "texto", "idioma", "sentimento_humano", "aspecto_humano", "observacao"])
    for r in A.itertuples():
        ws.append([r.id_amostra, r.text, r.lang, None, None, None])
    n = len(A) + 1
    dv1 = DataValidation(type="list", formula1='"negativo,neutro,positivo"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(list(cfg.ASPECT_SEEDS) + [GENERAL_BUCKET, "nenhum"]) + '"', allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"D2:D{n}")
    dv2.add(f"E2:E{n}")
    ws.column_dimensions["B"].width = 100
    ws.freeze_panes = "A2"
    wg = wb.create_sheet("Gabarito_modelo")
    wg.append(["id_amostra", "sent_id", "topico", "aspecto_modelo", "estrelas_modelo", "p_neg", "p_neu", "p_pos"])
    for r in A.itertuples():
        wg.append([r.id_amostra, r.sent_id, r.topic, r.aspecto, round(float(r.stars), 3), round(float(r.p_neg), 3),
                   round(float(r.p_neu), 3), round(float(r.p_pos), 3)])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log.warning("Planilha de validação manual criada em %s (%d frases). Rotule 'Amostra' e rode a etapa 14 de novo.", path, len(A))


def avaliar_amostra_manual(cfg, store: Store, path: Path) -> bool:
    from sklearn.metrics import classification_report, cohen_kappa_score, confusion_matrix, f1_score

    H = pd.read_excel(path, sheet_name="Amostra")
    G = pd.read_excel(path, sheet_name="Gabarito_modelo")
    H = H.merge(G, on="id_amostra")
    H = H[H["sentimento_humano"].notna() & H["aspecto_humano"].notna()].copy()
    if len(H) < 20:
        log.info("Amostra manual: %d frases rotuladas (< 20); métricas não calculadas.", len(H))
        return False
    if store.suffix:  # modelo B: as probabilidades/estrelas/classe vêm da tabela de B (a planilha de rótulos guarda as de A, arredondadas)
        from .sentimento import load_sentiment
        M = load_sentiment(store)[["text_id", "stars", "p_neg", "p_neu", "p_pos", "cls"]].rename(columns={"stars": "estrelas_modelo"})
        H = H.drop(columns=["estrelas_modelo", "p_neg", "p_neu", "p_pos"]).merge(store.load("sentences")[["sent_id", "text_id"]], on="sent_id").merge(M, on="text_id")
        H["sent_modelo"] = np.array(["negativo", "neutro", "positivo"])[H["cls"].to_numpy()]
    else:
        H["sent_modelo"] = np.array(["negativo", "neutro", "positivo"])[H[["p_neg", "p_neu", "p_pos"]].to_numpy().argmax(1)]
    H["asp_humano"] = H["aspecto_humano"].replace({"nenhum": GENERAL_BUCKET})
    labs = ["negativo", "neutro", "positivo"]
    num = {"negativo": -1, "neutro": 0, "positivo": 1}
    rows = [{"medida": "n frases rotuladas", "valor": len(H)},
            {"medida": "sentimento: acurácia", "valor": float((H["sent_modelo"] == H["sentimento_humano"]).mean())},
            {"medida": "sentimento: F1 macro", "valor": float(f1_score(H["sentimento_humano"], H["sent_modelo"], average="macro", labels=labs))},
            {"medida": "sentimento: kappa de Cohen", "valor": float(cohen_kappa_score(H["sentimento_humano"], H["sent_modelo"]))},
            {"medida": "sentimento: kappa ponderado linear", "valor": float(cohen_kappa_score(
                H["sentimento_humano"].map(num), H["sent_modelo"].map(num), weights="linear"))},
            {"medida": "sentimento: Spearman (humano -1/0/1 × estrelas do modelo)",
             "valor": float(stats.spearmanr(H["sentimento_humano"].map(num), H["estrelas_modelo"])[0])},
            {"medida": "aspecto: acurácia", "valor": float((H["aspecto_modelo"] == H["asp_humano"]).mean())},
            {"medida": "aspecto: kappa de Cohen", "valor": float(cohen_kappa_score(H["asp_humano"], H["aspecto_modelo"]))}]
    for lang, g in H.groupby("idioma"):
        if len(g) >= 10:
            rows.append({"medida": f"sentimento: acurácia [{lang}] (n={len(g)})", "valor": float((g["sent_modelo"] == g["sentimento_humano"]).mean())})
    store.save("manual_resumo", pd.DataFrame(rows))
    store.save("manual_confusao_sentimento", pd.crosstab(H["sentimento_humano"], H["sent_modelo"]).reset_index())
    store.save("manual_confusao_aspecto", pd.crosstab(H["asp_humano"], H["aspecto_modelo"]).reset_index())
    rep = pd.DataFrame(classification_report(H["asp_humano"], H["aspecto_modelo"], zero_division=0, output_dict=True)).T.reset_index()
    store.save("manual_aspecto_por_classe", rep.rename(columns={"index": "aspecto"}))
    log.info("Validação manual:\n%s", pd.DataFrame(rows).round(3).to_string(index=False))
    store.metric("etapa14_manual", n=len(H), acuracia_sentimento=round(rows[1]["valor"], 3), kappa_sentimento=round(rows[3]["valor"], 3),
                 acuracia_aspecto=round(rows[6]["valor"], 3))
    return True


def run(cfg, store: Store) -> None:
    with stage("Etapa 14 — Validade de construto ampliada"):
        D = store.load("perfil_absa").merge(store.load("listings_elig"), on="id")
        V = por_volume(cfg, D)
        M, R = matriz(cfg, D)
        A = auc_kendall(cfg, D)
        C = confiabilidade(cfg, store, D)
        Rc = recortes(cfg, store, D)
        for nome, t in [("val_volume", V), ("val_matriz", M), ("val_convergente", R), ("val_auc_kendall", A),
                        ("val_confiabilidade", C), ("val_recortes", Rc)]:
            store.save(nome, t)
        log.info("(a) Concordância por faixa de volume (Spearman):\n%s",
                 V.pivot_table(index="aspecto", columns="faixa", values="spearman").round(3).to_string())
        log.info("(b) Convergente/discriminante:\n%s", R.round(3).to_string(index=False))
        log.info("(c) Confiabilidade e correção da atenuação:\n%s", C.round(3).to_string(index=False))
        log.info("(d) Subnotas saturadas (AUC, Kendall):\n%s", A.round(3).to_string(index=False))
        log.info("(e) Recortes:\n%s", Rc.pivot_table(index="aspecto", columns="recorte", values="spearman").round(3).to_string())
        path = Path(cfg.MANUAL_VALIDATION_XLSX)
        if not path.exists():
            gerar_amostra_manual(cfg, store, path)
        else:
            avaliar_amostra_manual(cfg, store, path)
        store.metric("etapa14", spearman_correspondente=dict(zip(R["aspecto"], R["spearman_correspondente"].round(3))))
